from flask import render_template, redirect, url_for, flash, request, send_from_directory
from werkzeug.utils import secure_filename
from app import db, main_bp
from app.models import Contract, Category, Attachment, ContractItem
from app.forms import ContractForm, CategoryForm, SearchForm, ImportContractForm
import os
from datetime import datetime, date
import uuid
import pandas as pd
from openpyxl import load_workbook

from flask import current_app as app

@main_bp.route('/')
def index():
    contracts = Contract.query.order_by(Contract.created_at.desc()).limit(10).all()
    return render_template('index.html', contracts=contracts)

@main_bp.route('/contracts')
def list_contracts():
    page = request.args.get('page', 1, type=int)
    search_form = SearchForm()
    
    # 预填充分类下拉列表
    categories = Category.query.order_by(Category.name).all()
    search_form.category.choices = [(0, '全部')] + [(c.id, c.name) for c in categories]
    
    query = Contract.query
    
    # 处理搜索筛选
    keyword = request.args.get('keyword', '')
    category_id = request.args.get('category', 0, type=int)
    status = request.args.get('status', '')
    
    if keyword:
        query = query.filter(
            (Contract.title.like(f'%{keyword}%')) |
            (Contract.number.like(f'%{keyword}%')) |
            (Contract.party_a.like(f'%{keyword}%')) |
            (Contract.party_b.like(f'%{keyword}%'))
        )
    
    if category_id > 0:
        query = query.filter_by(category_id=category_id)
        
    if status:
        query = query.filter_by(status=status)
    
    contracts = query.order_by(Contract.updated_at.desc()).paginate(page=page, per_page=20)
    
    return render_template('contracts/list.html', 
                         contracts=contracts, 
                         form=search_form,
                         categories=categories)

@main_bp.route('/contracts/new', methods=['GET', 'POST'])
def new_contract():
    form = ContractForm()
    form.category_id.choices = [(c.id, c.name) for c in Category.query.all()]
    
    if form.validate_on_submit():
        contract = Contract(
            title=form.title.data,
            number=form.number.data,
            party_a=form.party_a.data,
            party_b=form.party_b.data,
            sign_date=form.sign_date.data,
            effective_date=form.effective_date.data,
            expiry_date=form.expiry_date.data,
            value=form.value.data,
            status=form.status.data,
            description=form.description.data,
            category_id=form.category_id.data
        )
        
        db.session.add(contract)
        db.session.commit()
        
        # 处理合同明细
        item_names = request.form.getlist('item-name')
        item_specs = request.form.getlist('item-spec')
        item_units = request.form.getlist('item-unit')
        item_quantities = request.form.getlist('item-quantity')
        item_prices = request.form.getlist('item-price')
        item_amounts = request.form.getlist('item-amount')
        
        for i in range(len(item_names)):
            item = ContractItem(
                name=item_names[i],
                specification=item_specs[i],
                unit=item_units[i],
                quantity=float(item_quantities[i]),
                unit_price=float(item_prices[i]),
                amount=float(item_amounts[i]),
                contract_id=contract.id
            )
            db.session.add(item)
        
        # 处理附件上传 - 支持多文件上传
        if form.attachments.data:
            attachments = save_attachment(form.attachments.data, contract.id)
            if not attachments:
                flash('附件上传失败或未选择有效文件', 'warning')
            
        db.session.commit()
        flash('合同创建成功!', 'success')
        return redirect(url_for('main.list_contracts'))
    
    return render_template('contracts/new.html', form=form)

@main_bp.route('/contracts/<int:id>/edit', methods=['GET', 'POST'])
def edit_contract(id):
    contract = Contract.query.get_or_404(id)
    form = ContractForm(obj=contract)
    form.category_id.choices = [(c.id, c.name) for c in Category.query.all()]
    
    # 预填充明细项数据到模板上下文
    items = contract.items
    
    if form.validate_on_submit():
        contract.title = form.title.data
        contract.number = form.number.data
        contract.party_a = form.party_a.data
        contract.party_b = form.party_b.data
        contract.sign_date = form.sign_date.data
        contract.effective_date = form.effective_date.data
        contract.expiry_date = form.expiry_date.data
        contract.value = form.value.data
        contract.status = form.status.data
        contract.description = form.description.data
        contract.category_id = form.category_id.data
        
        db.session.commit()
        
        # 处理附件上传 - 支持多文件上传
        if form.attachments.data:
            attachments = save_attachment(form.attachments.data, contract.id)
            if not attachments:
                flash('附件上传失败或未选择有效文件', 'warning')
            
        flash('合同更新成功!', 'success')
        return redirect(url_for('main.view_contract', id=contract.id))
    
    return render_template('contracts/edit.html', 
                         form=form, 
                         contract=contract,
                         items=items)

@main_bp.route('/contracts/<int:id>')
def view_contract(id):
    contract = Contract.query.get_or_404(id)
    return render_template('contracts/view.html', 
                         contract=contract,
                         items=contract.items)

@main_bp.route('/contracts/<int:id>/delete', methods=['POST'])
def delete_contract(id):
    contract = Contract.query.get_or_404(id)
    
    # 删除相关明细项
    for item in contract.items:
        db.session.delete(item)
        
    # 删除相关附件
    for attachment in contract.attachments:
        try:
            os.remove(os.path.join(app.config['UPLOAD_FOLDER'], attachment.filepath))
        except:
            pass
        db.session.delete(attachment)
    
    db.session.delete(contract)
    db.session.commit()
    flash('合同删除成功!', 'success')
    return redirect(url_for('main.list_contracts'))

@main_bp.route('/categories', methods=['GET', 'POST'])
def manage_categories():
    form = CategoryForm()
    categories = Category.query.all()
    
    if form.validate_on_submit():
        category = Category(
            name=form.name.data,
            description=form.description.data
        )
        db.session.add(category)
        db.session.commit()
        flash('分类添加成功!', 'success')
        return redirect(url_for('main.manage_categories'))
        
    return render_template('categories/manage.html', categories=categories, form=form)

def save_attachment(files, contract_id):
    saved_attachments = []
    try:
        # 确保上传目录存在
        os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
        
        # 处理单个文件或文件列表
        file_list = files if isinstance(files, list) else [files]
        
        for file in file_list:
            # 跳过无效文件对象
            if not file or not hasattr(file, 'filename'):
                continue
                
            filename = secure_filename(file.filename)
            if not filename:
                continue
                
            # 保留原始文件扩展名
            file_ext = os.path.splitext(filename)[1]
            unique_filename = f"{uuid.uuid4()}{file_ext}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            
            # 只处理 Werkzeug FileStorage 对象
            if hasattr(file, 'save'):
                file.save(filepath)
            else:
                # 跳过非文件上传对象
                continue
            
            # 验证文件是否保存成功
            if not os.path.exists(filepath):
                continue
                
            # 获取文件大小
            filesize = os.path.getsize(filepath)
            
            # 创建附件记录
            attachment = Attachment(
                filename=filename,
                filepath=unique_filename,
                mimetype=file.content_type,
                filesize=filesize,
                contract_id=contract_id
            )
            
            db.session.add(attachment)
            saved_attachments.append(attachment)
            
            app.logger.info(f"成功保存附件: {filename} (大小: {filesize}字节)")
        
        db.session.commit()
        return saved_attachments
        
    except Exception as e:
        app.logger.error(f"保存附件失败: {str(e)}")
        if 'filepath' in locals() and os.path.exists(filepath):
            os.remove(filepath)
        raise

@main_bp.route('/attachments/<int:id>')
def download_attachment(id):
    attachment = Attachment.query.get_or_404(id)
    try:
        # 确保文件名包含正确扩展名
        filename = attachment.filename
        if not any(filename.lower().endswith(ext) for ext in ['.pdf','.doc','.docx','.xls','.xlsx','.jpg','.png']):
            # 从mimetype推断扩展名
            ext_map = {
                'application/pdf': '.pdf',
                'application/msword': '.doc',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document': '.docx',
                'application/vnd.ms-excel': '.xls',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': '.xlsx',
                'image/jpeg': '.jpg',
                'image/png': '.png'
            }
            ext = ext_map.get(attachment.mimetype, '')
            filename = f"{filename}{ext}"
        
        # 确保下载文件名包含正确扩展名
        if not filename.lower().endswith(os.path.splitext(attachment.filename)[1].lower()):
            filename = f"{filename}{os.path.splitext(attachment.filename)[1]}"
        
        response = send_from_directory(
            app.config['UPLOAD_FOLDER'],
            attachment.filepath,
            as_attachment=True,
            download_name=filename,
            mimetype=attachment.mimetype
        )
        
        # 设置Content-Disposition头确保浏览器正确处理文件名
        response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        app.logger.error(f"下载附件失败: {str(e)}")
        flash('附件下载失败，文件可能已被删除', 'danger')
        return redirect(url_for('main.view_contract', id=attachment.contract_id))

@main_bp.route('/attachments/<int:id>/delete', methods=['POST'])
def delete_attachment(id):
    attachment = Attachment.query.get_or_404(id)
    contract_id = attachment.contract_id
    
    # 删除文件
    try:
        os.remove(os.path.join(app.config['UPLOAD_FOLDER'], attachment.filepath))
    except:
        pass
    
    # 删除数据库记录
    db.session.delete(attachment)
    db.session.commit()
    
    flash('附件删除成功!', 'success')
    return redirect(url_for('main.view_contract', id=contract_id))

@main_bp.route('/categories/<int:id>/edit', methods=['GET', 'POST'])
def edit_category(id):
    category = Category.query.get_or_404(id)
    form = CategoryForm(obj=category)
    
    if form.validate_on_submit():
        category.name = form.name.data
        category.description = form.description.data
        db.session.commit()
        flash('分类更新成功!', 'success')
        return redirect(url_for('main.manage_categories'))
    
    return render_template('categories/edit.html', form=form, category=category)

@main_bp.route('/contracts/import', methods=['GET', 'POST'])
def import_contracts():
    form = ImportContractForm()
    if form.validate_on_submit():
        try:
            # 保存上传的Excel文件
            filename = secure_filename(form.excel_file.data.filename)
            if not filename.lower().endswith(('.xlsx', '.xls')):
                flash('请上传有效的Excel文件(.xlsx或.xls格式)', 'danger')
                return redirect(url_for('main.import_contracts'))
                
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], f"import_{uuid.uuid4()}_{filename}")
            form.excel_file.data.save(filepath)
            
            try:
                # 验证并读取Excel文件
                wb = load_workbook(filepath)
                sheet = wb.active
                
                # 解析数据
                contracts_created = 0
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):  # 跳过空行
                        continue
                        
                    # 检查行数据长度(模板有8列)
                    if len(row) < 8:
                        raise ValueError(f"Excel第{row_idx}行数据不完整，需要8列(标题、编号、甲方、乙方、签订日期、生效日期、到期日期、金额)，当前行只有{len(row)}列")

                    # 检查必填字段
                    required_fields = [row[0], row[1], row[2], row[3]]
                    if not all(required_fields):
                        raise ValueError(f"Excel第{row_idx}行: 合同标题、编号、甲方、乙方为必填字段")

                    # 处理日期字段
                    def parse_date(date_input):
                        if isinstance(date_input, datetime):
                            return date_input.date()
                        if isinstance(date_input, date):
                            return date_input
                        try:
                            if isinstance(date_input, str):
                                return datetime.strptime(date_input, '%Y-%m-%d').date()
                            if isinstance(date_input, (int, float)):
                                return datetime.fromtimestamp(date_input).date()
                        except:
                            return datetime.now().date()
                        return datetime.now().date()

                    # 创建合同
                    contract = Contract(
                        title=str(row[0]).strip(),
                        number=str(row[1]).strip(),
                        party_a=str(row[2]).strip(),
                        party_b=str(row[3]).strip(),
                        sign_date=parse_date(row[4]),
                        effective_date=parse_date(row[5]),
                        expiry_date=parse_date(row[6]),
                        value=float(row[7]) if row[7] else 0,
                        status='active',  # 默认状态
                        description='',   # 默认空描述
                        category_id=None,  # 默认无分类
                        created_at=datetime.now(),
                        updated_at=datetime.now()
                    )
                    db.session.add(contract)
                    db.session.flush()  # 获取合同ID
                    
                    # 处理明细项
                    if len(row) > 11 and row[11]:  # 如果有明细项
                        try:
                            items = eval(row[11])  # 假设明细项以列表形式存储
                            for item_data in items:
                                item = ContractItem(
                                    name=item_data.get('name'),
                                    specification=item_data.get('spec'),
                                    unit=item_data.get('unit'),
                                    quantity=float(item_data.get('quantity', 0)),
                                    unit_price=float(item_data.get('unit_price', 0)),
                                    amount=float(item_data.get('amount', 0)),
                                    contract_id=contract.id
                                )
                                db.session.add(item)
                        except Exception as e:
                            raise ValueError(f"Excel第{row_idx}行: 明细项格式错误 - {str(e)}")
                
                    contracts_created += 1
                    app.logger.info(f"成功导入合同: {contract.title} (编号: {contract.number})")
            except Exception as e:
                db.session.rollback()
                flash(f'Excel文件处理失败: {str(e)}', 'danger')
                return redirect(url_for('main.import_contracts'))
            
            db.session.commit()
            flash(f'成功导入 {contracts_created} 份合同!', 'success')
            return redirect(url_for('main.list_contracts'))
            
        except Exception as e:
            db.session.rollback()
            flash(f'导入失败: {str(e)}', 'danger')
            return redirect(url_for('main.import_contracts'))
    
    return render_template('contracts/import.html', form=form)

@main_bp.route('/contracts/template')
def download_template():
    """下载合同导入模板"""
    return send_from_directory(
        os.path.join(app.root_path, 'static', 'templates'),
        'contract_import_template.xlsx',
        as_attachment=True,
        download_name='合同导入模板.xlsx'
    )

@main_bp.route('/categories/<int:id>/delete', methods=['POST'])
def delete_category(id):
    category = Category.query.get_or_404(id)
    
    # 检查是否有关联合同
    if category.contracts:
        flash('该分类下存在合同，无法删除！', 'danger')
        return redirect(url_for('main.manage_categories'))
    
    db.session.delete(category)
    db.session.commit()
    flash('分类删除成功！', 'success')
    return redirect(url_for('main.manage_categories'))
